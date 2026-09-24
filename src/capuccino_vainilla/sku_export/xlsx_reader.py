"""Lectura de SKUs desde una planilla .xlsx multi-hoja.

Cada hoja puede tener su propia fila de encabezado (varía según el
proveedor/marca), así que no se asume una posición fija: se busca, en las
primeras filas de cada hoja, una celda cuyo texto normalizado coincida con
``column_name`` y se leen los valores no vacíos debajo de esa columna.
"""

from __future__ import annotations

from openpyxl import load_workbook

from ..exceptions import SkuExportError

# Cuántas filas iniciales de cada hoja se inspeccionan buscando el encabezado.
# Las hojas de este workbook tienen entre 4 y 6 filas de metadata (proveedor,
# tipo de cambio, descuentos) antes del encabezado real.
_HEADER_SEARCH_ROWS = 15


def _normalize(text: object) -> str:
    return str(text).strip().upper() if text is not None else ""


def _find_header_cell(sheet, column_name: str) -> tuple[int, int] | None:
    """Devuelve ``(fila, columna)`` del encabezado, o None si no está en la hoja."""
    target = _normalize(column_name)
    max_row = min(_HEADER_SEARCH_ROWS, sheet.max_row or 0)
    for row in range(1, max_row + 1):
        for col in range(1, (sheet.max_column or 0) + 1):
            if _normalize(sheet.cell(row=row, column=col).value) == target:
                return row, col
    return None


def extract_skus(path: str, column_name: str = "REF ODOO") -> frozenset[str]:
    """Lee ``column_name`` en cada hoja de ``path`` y devuelve el conjunto de SKUs.

    Hojas sin esa columna se omiten (no todas las hojas necesitan tenerla).
    Falla si ninguna hoja la tiene, o si el resultado queda vacío: un archivo
    de lista blanca vacío bloquearía TODA la sincronización de catálogo
    (ver ``config._load_sku_allowlist``), así que es mejor fallar temprano acá.
    """
    try:
        # read_only=False (default): calcula max_row/max_column recorriendo
        # las celdas reales en vez de confiar en la etiqueta <dimension> del
        # XML, que puede quedar desactualizada si algo insertó columnas sin
        # tocarla (ej. una edición quirúrgica del XML fuera de Excel).
        wb = load_workbook(path, data_only=True)
    except (OSError, KeyError) as exc:
        raise SkuExportError(f"No se pudo abrir '{path}': {exc}") from exc

    skus: set[str] = set()
    sheets_with_column = 0
    try:
        for sheet in wb.worksheets:
            header = _find_header_cell(sheet, column_name)
            if header is None:
                continue
            sheets_with_column += 1
            header_row, col = header
            for row in range(header_row + 1, (sheet.max_row or header_row) + 1):
                value = sheet.cell(row=row, column=col).value
                text = _normalize(value)
                if text:
                    skus.add(text)
    finally:
        wb.close()

    if sheets_with_column == 0:
        raise SkuExportError(
            f"Ninguna hoja de '{path}' tiene una columna '{column_name}'."
        )
    if not skus:
        raise SkuExportError(
            f"La columna '{column_name}' de '{path}' no tiene ningún valor cargado."
        )
    return frozenset(skus)
