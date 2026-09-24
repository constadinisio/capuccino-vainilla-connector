"""Tests de extracción de SKUs desde una planilla .xlsx multi-hoja."""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from capuccino_vainilla.exceptions import SkuExportError
from capuccino_vainilla.sku_export.xlsx_reader import extract_skus


def _write_sheet(wb, name, header_row, header_col, header_text, rows):
    """Crea una hoja con el encabezado en una posición arbitraria (simula el
    layout real, donde cada hoja de marca tiene su propia fila de metadata)."""
    ws = wb.create_sheet(name)
    ws.cell(row=header_row, column=header_col, value=header_text)
    for i, value in enumerate(rows, start=1):
        ws.cell(row=header_row + i, column=header_col, value=value)
    return ws


def test_extracts_skus_across_sheets_with_different_header_rows(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "AIDA", header_row=6, header_col=14, header_text="REF ODOO",
                 rows=["CAMAID001", "CAMAID009", None])
    _write_sheet(wb, "MEIKE", header_row=7, header_col=9, header_text="REF ODOO",
                 rows=["LENMEI003", "LENMEI004"])
    path = tmp_path / "planilla.xlsx"
    wb.save(path)

    skus = extract_skus(str(path))
    assert skus == frozenset({"CAMAID001", "CAMAID009", "LENMEI003", "LENMEI004"})


def test_sheet_without_the_column_is_skipped(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "CONMATCH", header_row=6, header_col=14, header_text="REF ODOO",
                 rows=["CAMAID001"])
    ws_no_match = wb.create_sheet("SINMATCH")
    ws_no_match.cell(row=6, column=1, value="MARCA")
    ws_no_match.cell(row=7, column=1, value="ALGO")
    path = tmp_path / "planilla.xlsx"
    wb.save(path)

    skus = extract_skus(str(path))
    assert skus == frozenset({"CAMAID001"})


def test_duplicates_collapse_to_one(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "S1", header_row=1, header_col=1, header_text="REF ODOO",
                 rows=["CAMAID001", "CAMAID001", "  CAMAID001  "])
    path = tmp_path / "planilla.xlsx"
    wb.save(path)

    skus = extract_skus(str(path))
    assert skus == frozenset({"CAMAID001"})


def test_header_matching_is_case_and_whitespace_insensitive(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "S1", header_row=1, header_col=1, header_text="  ref odoo  ",
                 rows=["CAMAID001"])
    path = tmp_path / "planilla.xlsx"
    wb.save(path)

    skus = extract_skus(str(path), column_name="REF ODOO")
    assert skus == frozenset({"CAMAID001"})


def test_custom_column_name(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "S1", header_row=1, header_col=1, header_text="SKU ODOO",
                 rows=["X-1"])
    path = tmp_path / "planilla.xlsx"
    wb.save(path)

    skus = extract_skus(str(path), column_name="SKU ODOO")
    assert skus == frozenset({"X-1"})


def test_no_sheet_has_the_column_raises(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("S1")
    ws.cell(row=1, column=1, value="MARCA")
    path = tmp_path / "planilla.xlsx"
    wb.save(path)

    with pytest.raises(SkuExportError, match="REF ODOO"):
        extract_skus(str(path))


def test_column_present_but_empty_raises(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "S1", header_row=1, header_col=1, header_text="REF ODOO",
                 rows=[None, "  ", None])
    path = tmp_path / "planilla.xlsx"
    wb.save(path)

    with pytest.raises(SkuExportError, match="ningún valor"):
        extract_skus(str(path))


def test_missing_file_raises(tmp_path):
    missing = tmp_path / "no-existe.xlsx"
    with pytest.raises(SkuExportError, match="no-existe.xlsx"):
        extract_skus(str(missing))
